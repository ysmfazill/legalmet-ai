"""Report DOCX renderer (UI-08 §15).

A REAL editable document built with python-docx from the FROZEN snapshot.
Same core information as the PDF: summary, source complaint, findings,
regulatory basis, physical verification, lot intelligence, decision,
traceability — professional formatting, editable in Word.

Safety: no fake government letterhead, no official seals, no credentials.
python-docx inserts plain text runs (no HTML passthrough), so user-controlled
strings cannot become executable content.
"""
from __future__ import annotations

import io
from typing import Any, Callable

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

NAVY = "0F2A43"
MUTED = "5A6B7A"

MAX_IMAGES = 12


def _fmt(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _add_kv_table(doc, rows: list[tuple[str, str]]) -> None:  # noqa: ANN001
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for key, val in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = val or "—"
        for run in cells[0].paragraphs[0].runs or []:
            run.bold = True
            run.font.size = Pt(8)
        for cell in cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs or []:
                    run.font.size = Pt(8)


def _add_grid_table(doc, header: list[str], rows: list[list[str]]) -> None:  # noqa: ANN001
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for cell, text in zip(table.rows[0].cells, header, strict=False):
        cell.text = text
        for paragraph in cell.paragraphs:
            for run in paragraph.runs or []:
                run.bold = True
                run.font.size = Pt(7.5)
    for row_values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_values, strict=False):
            cell.text = text or "—"
            for paragraph in cell.paragraphs:
                for run in paragraph.runs or []:
                    run.font.size = Pt(7.5)


def _add_note(doc, text: str) -> None:  # noqa: ANN001
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = _rgb(MUTED)


def _rgb(hex_str: str):  # noqa: ANN202
    from docx.shared import RGBColor

    return RGBColor.from_string(hex_str)


def _add_heading(doc, text: str) -> None:  # noqa: ANN001
    heading = doc.add_heading(text, level=2)
    for run in heading.runs or []:
        run.font.color.rgb = _rgb(NAVY)
        run.font.size = Pt(11)


def render_report_docx(
    snapshot: dict, *, evidence_items: list[dict], image_loader, audit_events: list[dict]
) -> bytes:
    """Build the editable DOCX (image_loader: detail -> bytes | None)."""
    doc = DocxDocument()
    inspection = snapshot.get("inspection") or {}
    complaint = snapshot.get("sourceComplaint")
    findings = snapshot.get("findings") or []
    measurements = snapshot.get("measurements") or []
    lots = snapshot.get("lots") or []
    decision = snapshot.get("decision")
    basis = snapshot.get("regulatoryBasis") or []

    # ---------------------------------------------------------------- header
    title = doc.add_heading("METRASIGHT Inspection Report", level=0)
    for run in title.runs or []:
        run.font.color.rgb = _rgb(NAVY)
    subtitle = doc.add_paragraph()
    run = subtitle.add_run(
        "AI-assisted Legal Metrology inspection — inspector-reviewed, "
        "evidence-backed decision-support artifact."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(MUTED)

    _add_kv_table(
        doc,
        [
            ("Report ID", str(snapshot.get("reportId", ""))[:36]),
            ("Inspection ID", str(snapshot.get("inspectionId", ""))[:36]),
            ("Inspection ref.", _fmt(inspection.get("reference"))),
            ("Report version", f"v{_fmt(snapshot.get('version', 1))}"),
            ("Result", _fmt(snapshot.get("result"))),
            ("Generated", _fmt(snapshot.get("generatedAt"))),
            ("Product", _fmt(inspection.get("productName"))),
            ("Inspector", _fmt(inspection.get("inspectorName")) or "NOT RECORDED"),
            ("Inspection date", _fmt(inspection.get("date"))),
            ("Status", _fmt(inspection.get("status"))),
        ],
    )
    _add_note(
        doc,
        "Reports are decision-support artifacts generated from inspection "
        "evidence. They do not replace the authority of the authorized Legal "
        "Metrology inspector.",
    )

    # -------------------------------------------------------- source context
    _add_heading(doc, "Source Context")
    if complaint:
        _add_note(
            doc,
            "This inspection originated from a citizen complaint. SOURCE "
            "evidence (citizen-submitted) is reported separately from OFFICIAL "
            "inspection evidence and is never merged.",
        )
        _add_kv_table(
            doc,
            [
                ("Complaint ref.", _fmt(complaint.get("reference"))),
                ("Status", _fmt(complaint.get("status"))),
                ("Product", _fmt(complaint.get("product"))),
                ("Location", _fmt(complaint.get("location"))),
                ("Issue", _fmt(complaint.get("issue"))),
                ("Priority", _fmt(complaint.get("priority"))),
            ],
        )
    else:
        doc.add_paragraph("Routine inspection — no citizen complaint origin.")

    # ------------------------------------------------------------- findings
    _add_heading(doc, "Executive Findings")
    if findings:
        _add_grid_table(
            doc,
            ["ID", "Status", "Severity", "Requirement", "Rule", "Evidence", "Review"],
            [
                [
                    f["id"][:8],
                    _fmt(f.get("status")),
                    _fmt(f.get("severity")),
                    _fmt(f.get("requirement"))[:90],
                    _fmt(f.get("ruleCode")),
                    _fmt(f.get("evidenceStatus")),
                    _fmt(f.get("reviewState")),
                ]
                for f in findings
            ],
        )
    else:
        doc.add_paragraph("No evaluation findings recorded — NOT EVALUATED.")

    # ------------------------------------------------------ regulatory basis
    _add_heading(doc, "Regulatory Basis")
    for entry in basis:
        if entry.get("status") == "NOT_EVALUATED" or not entry.get("engineVersion"):
            doc.add_paragraph("NOT EVALUATED — REGULATORY REVIEW REQUIRED")
            continue
        _add_kv_table(
            doc,
            [
                ("Engine version", _fmt(entry.get("engineVersion"))),
                ("Regulatory version", _fmt(entry.get("regulatoryVersionLabel"))),
                ("Context date", _fmt(entry.get("contextDate"))),
            ],
        )
    _add_note(
        doc,
        "Regulatory basis is quoted from the deterministic engine output "
        "frozen at evaluation time. Unverified research-grade data is never "
        "presented as validated legal advice.",
    )

    # -------------------------------------------------- physical verification
    _add_heading(doc, "Physical Verification")
    if measurements:
        _add_grid_table(
            doc,
            ["Declared", "Measured", "Observed diff.", "Instrument", "Recorded", "Evaluation"],
            [
                [
                    _fmt(m.get("declaredValue")),
                    f"{_fmt(m.get('measuredValue'))} {_fmt(m.get('unit'))}".strip(),
                    _fmt(m.get("observedDifference")),
                    _fmt(m.get("instrumentId")) or "manual entry",
                    _fmt(m.get("recordedAt"))[:19],
                    _fmt((m.get("evaluation") or {}).get("status")),
                ]
                for m in measurements
            ],
        )
        _add_note(
            doc,
            "The observed difference is arithmetic on normalized values — a "
            "legal conclusion appears only from the frozen regulatory "
            "evaluation column, never from the difference itself.",
        )
    else:
        doc.add_paragraph("No physical measurements recorded.")

    # -------------------------------------------------------- lot intelligence
    if lots:
        _add_heading(doc, "Lot Intelligence")
        _add_grid_table(
            doc,
            ["Lot", "Declared", "Lot size", "Sampled", "Measured", "Result"],
            [
                [
                    _fmt(l.get("label")),
                    _fmt(l.get("declaredValue")),
                    _fmt(l.get("lotSize")),
                    _fmt(l.get("sampled")),
                    _fmt(l.get("measured")),
                    _fmt(l.get("decision")) or "IN PROGRESS",
                ]
                for l in lots
            ],
        )
        _add_note(
            doc,
            "Sampling shown is an AI-assisted inspection recommendation when "
            "no statutory sampling procedure applies — never a claim of the "
            "legally required sample.",
        )

    # -------------------------------------------------------- evidence images
    _add_heading(doc, "Evidence")
    # UI-10 §24: evidence_pack returns snake_case keys (evidence_type); accept
    # camelCase too. Before this fix evidence images were silently dropped.
    image_items = [
        i
        for i in evidence_items
        if (i.get("evidence_type") or i.get("evidenceType")) == "IMAGE"
    ]
    embedded = 0
    for item in image_items[:MAX_IMAGES]:
        detail = item.get("detail") or {}
        raw = image_loader(detail)
        optimized = _optimize(raw) if raw else None
        if optimized is None:
            _add_note(doc, f"{item.get('ref', '')} — Evidence image unavailable.")
            continue
        paragraph = doc.add_paragraph()
        caption = paragraph.add_run(
            f"{item.get('ref', '')} — Evidence image {_fmt(detail.get('filename'))}"
        )
        caption.font.size = Pt(8)
        caption.italic = True
        buf = io.BytesIO(optimized)
        run = paragraph.add_run()
        run.add_picture(buf, width=Mm(60))
        embedded += 1
    if len(image_items) > MAX_IMAGES:
        _add_note(
            doc,
            f"{len(image_items) - MAX_IMAGES} additional evidence images not "
            "embedded (export cap).",
        )

    _add_heading(doc, "Evidence Manifest")
    if evidence_items:
        _add_grid_table(
            doc,
            ["Ref", "Type", "Label", "Origin"],
            [
                [
                    _fmt(i.get("ref")),
                    _fmt(i.get("evidence_type") or i.get("evidenceType")),
                    _fmt(i.get("label"))[:80],
                    _fmt((i.get("detail") or {}).get("origin")),
                ]
                for i in evidence_items
            ],
        )
    else:
        doc.add_paragraph("No evidence manifest entries.")

    # ----------------------------------------------------- inspector decision
    _add_heading(doc, "Inspector Decision")
    if decision:
        _add_kv_table(
            doc,
            [
                ("Decision", _fmt(decision.get("decision"))),
                ("Reason", _fmt(decision.get("reason"))),
                ("Decided by", _fmt(decision.get("decidedByName"))),
                ("Decided at", _fmt(decision.get("decidedAt"))),
            ],
        )
    else:
        doc.add_paragraph(
            "NOT EVALUATED — no inspector decision recorded. The report "
            "cannot state a result without one."
        )
    _add_note(
        doc,
        "AI-assisted assessment. Final decision recorded by authorized "
        "inspector. The AI is never the final legal authority.",
    )

    # ------------------------------------------------------------ traceability
    _add_heading(doc, "Traceability")
    if audit_events:
        _add_grid_table(
            doc,
            ["Event", "Role", "Actor", "At"],
            [
                [
                    _fmt(e.get("eventType")),
                    _fmt(e.get("actorRole")),
                    _fmt(e.get("actorName")),
                    _fmt(e.get("createdAt"))[:19],
                ]
                for e in audit_events
            ],
        )
    _add_note(
        doc,
        f"Report version v{_fmt(snapshot.get('version', 1))} — snapshot frozen "
        "at generation time. Append-only audit trail preserved.",
    )

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _optimize(data: bytes) -> bytes | None:
    from app.services.report.evidence_media import optimize_image_bytes

    return optimize_image_bytes(data)
